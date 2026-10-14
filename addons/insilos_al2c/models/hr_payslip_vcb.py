# -*- coding: utf-8 -*-
# Part of Insilos AL2C Enterprise Architecture (SOP-DGM-HR-04 & VCB DigiBiz Payroll)

import base64
import io
import re
import unicodedata
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from odoo import api, fields, models, _
from odoo.exceptions import UserError

VCB_DIGIBIZ_COLUMNS = [
    'stt\n(No)',
    'so ref \n(ref no)',
    'so tai khoan\n(beneficiary account)',
    'so cmnd\n(ID number)',
    'ngay cap\n(issued date)',
    'noi cap\n(issued place)',
    'ten nguoi huong\n(beneficiary name)',
    'ten ngan hang\n(bank name)',
    'so tien\n(amount)',
    'loai tien\n(currency)',
    'noi dung\n(content)',
]


def remove_vietnamese_accents(text):
    """Normalize Vietnamese text into uppercase unaccented ASCII."""
    if not text:
        return ""
    text = str(text).strip()
    text = text.replace('đ', 'd').replace('Đ', 'D')
    text = unicodedata.normalize('NFD', text)
    text = ''.join(c for c in text if unicodedata.category(c) != 'Mn')
    return text.upper().strip()


def build_vcb_digibiz_workbook(payslip_records, month=None, year=None):
    """Core builder that compiles a set of hr.payslip records into the exact
    11-column Excel file conforming to bangkemau-VCB-DigiBiz (2).xlsx.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Write headers
    for col_idx, col_name in enumerate(VCB_DIGIBIZ_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.alignment = header_alignment

    # Determine period month and year
    if not month or not year:
        first_slip = payslip_records[:1]
        target_date = first_slip.date_to or first_slip.date_from or fields.Date.today()
        month = month or target_date.month
        year = year or target_date.year

    transfer_content = f"cong ty cp nguoi ban so thanh toan luong thang {int(month):02d} {int(year)}"

    row_idx = 2
    for seq, payslip in enumerate(payslip_records, start=1):
        emp = payslip.employee_id
        if not emp:
            continue

        # 1. STT
        ws.cell(row=row_idx, column=1, value=seq)

        # 2. So Ref
        ref_no = emp.registration_number or payslip.number or payslip.code or ""
        ws.cell(row=row_idx, column=2, value=str(ref_no) if ref_no else None)

        # 3. So tai khoan (Preserve leading zero as text)
        bank_acc = ""
        bank_record = False
        if hasattr(emp, 'primary_bank_account_id') and emp.primary_bank_account_id:
            bank_record = emp.primary_bank_account_id
            bank_acc = getattr(bank_record, 'account_number', None) or getattr(bank_record, 'acc_number', None) or ""
        elif hasattr(emp, 'bank_account_ids') and emp.bank_account_ids:
            bank_record = emp.bank_account_ids[0]
            bank_acc = getattr(bank_record, 'account_number', None) or getattr(bank_record, 'acc_number', None) or ""
        elif hasattr(emp, 'bank_account_id') and emp.bank_account_id:
            bank_record = emp.bank_account_id
            bank_acc = getattr(bank_record, 'account_number', None) or getattr(bank_record, 'acc_number', None) or ""
        elif hasattr(emp, 'work_contact_id') and emp.work_contact_id and getattr(emp.work_contact_id, 'bank_ids', False):
            bank_record = emp.work_contact_id.bank_ids[0]
            bank_acc = getattr(bank_record, 'account_number', None) or getattr(bank_record, 'acc_number', None) or ""

        clean_acc = re.sub(r'[^0-9A-Za-z]', '', str(bank_acc or '')).strip()
        acc_cell = ws.cell(row=row_idx, column=3, value=clean_acc if clean_acc else "")
        acc_cell.number_format = '@'  # Explicit text format

        # 4. So CMND/CCCD
        id_number = ""
        if hasattr(emp, 'identification_id') and emp.identification_id:
            id_number = emp.identification_id
        elif hasattr(emp, 'current_version_id') and getattr(emp.current_version_id, 'identification_id', False):
            id_number = emp.current_version_id.identification_id
        elif hasattr(emp, 'id_card_name') and emp.id_card_name:
            id_number = emp.id_card_name
        clean_id = re.sub(r'[^0-9A-Za-z]', '', str(id_number or '')).strip()
        ws.cell(row=row_idx, column=4, value=clean_id if clean_id else None)

        # 5. Ngay cap
        ws.cell(row=row_idx, column=5, value=None)

        # 6. Noi cap
        ws.cell(row=row_idx, column=6, value=None)

        # 7. Ten nguoi huong (UPPERCASE UNACCENTED)
        raw_name = emp.legal_name or emp.name or ""
        std_beneficiary = remove_vietnamese_accents(raw_name)
        ws.cell(row=row_idx, column=7, value=std_beneficiary)

        # 8. Ten ngan hang
        bank_name = "VIETCOMBANK"
        raw_b_name = ""
        if bank_record:
            if getattr(bank_record, 'bank_name', False):
                raw_b_name = bank_record.bank_name
            elif getattr(bank_record, 'bank_id', False) and bank_record.bank_id.name:
                raw_b_name = bank_record.bank_id.name

        if raw_b_name:
            b_name = raw_b_name.upper()
            if "ACB" in b_name or "A CHAU" in b_name:
                bank_name = "ACB"
            elif "VIETCOM" in b_name or "VCB" in b_name:
                bank_name = "VIETCOMBANK"
            else:
                bank_name = remove_vietnamese_accents(raw_b_name)
        ws.cell(row=row_idx, column=8, value=bank_name)

        # 9. So tien (Net Salary positive integer)
        net_amount = 0
        if getattr(payslip, 'al2c_net_salary', 0):
            net_amount = int(round(payslip.al2c_net_salary))
        elif hasattr(payslip, 'net_wage') and payslip.net_wage:
            net_amount = int(round(payslip.net_wage))
        elif hasattr(payslip, 'line_ids') and payslip.line_ids:
            net_line = payslip.line_ids.filtered(lambda l: l.code == 'NET')
            if net_line:
                net_amount = int(round(sum(net_line.mapped('total'))))
            else:
                net_amount = int(round(getattr(payslip, 'wage', 0) or 0))
        else:
            net_amount = int(round(getattr(payslip, 'wage', 0) or 0))

        ws.cell(row=row_idx, column=9, value=str(net_amount) if net_amount else "0")

        # 10. Loai tien
        ws.cell(row=row_idx, column=10, value="vnd")

        # 11. Noi dung
        ws.cell(row=row_idx, column=11, value=transfer_content)

        row_idx += 1

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


class HrPayslipRun(models.Model):
    _inherit = 'hr.payslip.run'

    def action_export_vcb_digibiz(self):
        """Action for UI button on Payslip Run to export VCB DigiBiz 11-column Excel."""
        self.ensure_one()
        excel_bytes, filename = self.generate_vcb_digibiz_xlsx()
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(excel_bytes),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'res_model': 'hr.payslip.run',
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

    def generate_vcb_digibiz_xlsx(self, payslips=None, month=None, year=None):
        """Generates 11-column Excel binary adhering to bangkemau-VCB-DigiBiz (2).xlsx.
        Returns: (bytes, filename)
        """
        self.ensure_one()
        slips = payslips or self.slip_ids.filtered(lambda p: p.state not in ('cancel',))
        if not slips:
            raise UserError(_("No active payslips available to export in pay run %s.", self.name))

        target_date = self.date_end or self.date_start or fields.Date.today()
        m = month or target_date.month
        y = year or target_date.year

        excel_data = build_vcb_digibiz_workbook(slips, month=m, year=y)
        filename = f"VCB_DigiBiz_Luong_Thang_{int(m):02d}_{int(y)}.xlsx"
        return excel_data, filename


class HrPayslip(models.Model):
    _inherit = 'hr.payslip'

    al2c_net_salary = fields.Monetary(
        string="Net Salary (AL2C)",
        help="Net salary amount for banking and payroll disbursements"
    )

    def _compute_basic_net(self):
        super()._compute_basic_net()
        for slip in self:
            if not slip.net_wage and getattr(slip, 'al2c_net_salary', 0):
                slip.net_wage = slip.al2c_net_salary

    def action_export_vcb_digibiz(self):
        """Action on payslip recordset to export VCB DigiBiz Excel."""
        excel_bytes, filename = self.generate_vcb_digibiz_xlsx()
        res_id = self[0].id if len(self) == 1 else False
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(excel_bytes),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'res_model': 'hr.payslip',
            'res_id': res_id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

    def generate_vcb_digibiz_xlsx(self, month=None, year=None):
        """Generates 11-column Excel binary for given payslips."""
        slips = self.filtered(lambda p: p.state not in ('cancel',))
        if not slips:
            raise UserError(_("No valid payslips selected for VCB DigiBiz export."))

        target_date = slips[0].date_to or slips[0].date_from or fields.Date.today()
        m = month or target_date.month
        y = year or target_date.year

        excel_data = build_vcb_digibiz_workbook(slips, month=m, year=y)
        filename = f"VCB_DigiBiz_Luong_Thang_{int(m):02d}_{int(y)}.xlsx"
        return excel_data, filename
