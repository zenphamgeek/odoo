# -*- coding: utf-8 -*-
# Part of Insilos AL2C Enterprise Architecture (DGM Governance & Hyper-Automation Suite)

import base64
import datetime
import io
import openpyxl
import re
from unittest.mock import patch

from odoo import fields
from odoo.tests.common import TransactionCase, tagged
from odoo.exceptions import UserError
from odoo.tools.pdf import PdfFileWriter

from odoo.addons.insilos_al2c.models.hr_payslip_vcb import VCB_DIGIBIZ_COLUMNS, remove_vietnamese_accents


@tagged('post_install', '-at_install', 'insilos_al2c', 'dgm_governance')
class TestDgmGovernanceAutomation(TransactionCase):
    """Holistic Unit & Integration Test Suite for DGM / Digimate JSC (MST: 0317958465)
    Hyper-Automation, Legal Entity Compliance Scanner, AI Sort Prompts, and SOP-RAG Indexing.
    """

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env.company
        cls.scanner = cls.env['al2c.compliance.scanner']
        cls.rag_service = cls.env['al2c.rag.service']
        cls.doc_model = cls.env['documents.document']

        # Setup Employees with diverse Vietnamese accented names and bank accounts
        cls.emp_hoang = cls.env['hr.employee'].create({
            'name': 'Lê Xuân Hoàng',
            'legal_name': 'Lê Xuân Hoàng',
            'registration_number': 'DGM-001',
            'company_id': cls.company.id,
        })
        cls.bank_hoang = cls.env['res.partner.bank'].create({
            'account_number': '1373068161',
            'bank_name': 'Vietcombank Nam Sai Gon',
            'bank_bic': 'BFTVVNVX',
            'partner_id': cls.emp_hoang.work_contact_id.id,
        })
        cls.emp_hoang.write({
            'bank_account_ids': [(4, cls.bank_hoang.id)],
            'primary_bank_account_id': cls.bank_hoang.id,
        })

        # Employee with leading zero account number
        cls.emp_tham = cls.env['hr.employee'].create({
            'name': 'Mai Thị Hồng Thắm',
            'legal_name': 'Mai Thị Hồng Thắm',
            'registration_number': 'DGM-003',
            'company_id': cls.company.id,
        })
        cls.bank_tham = cls.env['res.partner.bank'].create({
            'account_number': '0111000287863',
            'bank_name': 'Vietcombank Chi nhanh Ba Dinh',
            'bank_bic': 'BFTVVNVX',
            'partner_id': cls.emp_tham.work_contact_id.id,
        })
        cls.emp_tham.write({
            'bank_account_ids': [(4, cls.bank_tham.id)],
            'primary_bank_account_id': cls.bank_tham.id,
        })

        # Employee with ACB account
        cls.emp_khanh = cls.env['hr.employee'].create({
            'name': 'Huỳnh Trần Quang Khánh',
            'legal_name': 'Huỳnh Trần Quang Khánh',
            'registration_number': 'DGM-005',
            'company_id': cls.company.id,
        })
        cls.bank_khanh = cls.env['res.partner.bank'].create({
            'account_number': '12154887',
            'bank_name': 'Asia Commercial Bank (ACB)',
            'bank_bic': 'ASCBVNVX',
            'partner_id': cls.emp_khanh.work_contact_id.id,
        })
        cls.emp_khanh.write({
            'bank_account_ids': [(4, cls.bank_khanh.id)],
            'primary_bank_account_id': cls.bank_khanh.id,
        })

        # Setup test pay run and payslips
        cls.pay_run = cls.env['hr.payslip.run'].create({
            'name': 'Bảng lương DGM Tháng 10/2026',
            'date_start': fields.Date.from_string('2026-10-01'),
            'date_end': fields.Date.from_string('2026-10-31'),
            'company_id': cls.company.id,
        })

        cls.slip_hoang = cls.env['hr.payslip'].create({
            'name': 'Phiếu lương Lê Xuân Hoàng - 10/2026',
            'employee_id': cls.emp_hoang.id,
            'payslip_run_id': cls.pay_run.id,
            'date_from': cls.pay_run.date_start,
            'date_to': cls.pay_run.date_end,
            'company_id': cls.company.id,
        })
        # Set net wage
        cls.slip_hoang.write({'net_wage': 21410000, 'al2c_net_salary': 21410000})

        cls.slip_tham = cls.env['hr.payslip'].create({
            'name': 'Phiếu lương Mai Thị Hồng Thắm - 10/2026',
            'employee_id': cls.emp_tham.id,
            'payslip_run_id': cls.pay_run.id,
            'date_from': cls.pay_run.date_start,
            'date_to': cls.pay_run.date_end,
            'company_id': cls.company.id,
        })
        cls.slip_tham.write({'net_wage': 13469000, 'al2c_net_salary': 13469000})

        cls.slip_khanh = cls.env['hr.payslip'].create({
            'name': 'Phiếu lương Huỳnh Trần Quang Khánh - 10/2026',
            'employee_id': cls.emp_khanh.id,
            'payslip_run_id': cls.pay_run.id,
            'date_from': cls.pay_run.date_start,
            'date_to': cls.pay_run.date_end,
            'company_id': cls.company.id,
        })
        cls.slip_khanh.write({'net_wage': 24876000, 'al2c_net_salary': 24876000})

    # =========================================================================
    # 1. VCB DIGIBIZ 11-COLUMN PAYROLL EXPORT TESTS
    # =========================================================================
    def test_01_vcb_digibiz_11_column_export_structure(self):
        """Verify generation of the exact 11-column Excel batch file conforming to
        bangkemau-VCB-DigiBiz (2).xlsx: column headers, uppercase unaccented names,
        leading zeros preserved on account numbers, and standardized transfer content.
        """
        # Test accent removal helper
        self.assertEqual(remove_vietnamese_accents("Lê Xuân Hoàng"), "LE XUAN HOANG")
        self.assertEqual(remove_vietnamese_accents("Huỳnh Trần Quang Khánh"), "HUYNH TRAN QUANG KHANH")
        self.assertEqual(remove_vietnamese_accents("Mai Thị Hồng Thắm"), "MAI THI HONG THAM")
        self.assertEqual(remove_vietnamese_accents("Đỗ Văn Phong"), "DO VAN PHONG")

        # Generate workbook binary from pay run
        excel_bytes, filename = self.pay_run.generate_vcb_digibiz_xlsx()
        self.assertTrue(excel_bytes, "Must generate valid Excel binary content")
        self.assertEqual(filename, "VCB_DigiBiz_Luong_Thang_10_2026.xlsx")

        # Inspect workbook with openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertIn("Sheet1", wb.sheetnames, "Workbook must contain Sheet1")
        ws = wb["Sheet1"]

        # 1. Verify 11 Column Headers
        headers = [ws.cell(row=1, column=col).value for col in range(1, 12)]
        self.assertEqual(len(headers), 11, "Must have exactly 11 columns")
        self.assertEqual(headers, VCB_DIGIBIZ_COLUMNS, "Column headers must match VCB DigiBiz standard exactly")

        # 2. Verify Row Content
        rows = list(ws.iter_rows(min_row=2, max_row=4, values_only=True))
        self.assertEqual(len(rows), 3, "Must have 3 data rows")

        # Row 1: Le Xuan Hoang
        r1 = rows[0]
        self.assertEqual(r1[0], 1, "STT must be 1")
        self.assertEqual(r1[1], "DGM-001", "Ref must match employee reference")
        self.assertEqual(r1[2], "1373068161", "Account number must match")
        self.assertEqual(r1[6], "LE XUAN HOANG", "Beneficiary name must be UPPERCASE UNACCENTED")
        self.assertEqual(r1[7], "VIETCOMBANK", "Bank name must be VIETCOMBANK")
        self.assertEqual(str(r1[8]), "21410000", "Amount must match Net Salary")
        self.assertEqual(r1[9], "vnd", "Currency must be vnd")
        self.assertEqual(r1[10], "cong ty cp nguoi ban so thanh toan luong thang 10 2026", "Standard narration format")

        # Row 2: Mai Thi Hong Tham (Leading zero preserved)
        r2 = rows[1]
        self.assertEqual(r2[0], 2, "STT must be 2")
        self.assertEqual(r2[2], "0111000287863", "Leading zero must be strictly preserved as text")
        self.assertTrue(r2[2].startswith("0"), "Account number must start with 0")
        self.assertEqual(r2[6], "MAI THI HONG THAM", "Beneficiary name must be UPPERCASE UNACCENTED")
        self.assertEqual(r2[7], "VIETCOMBANK")
        self.assertEqual(str(r2[8]), "13469000")
        self.assertEqual(r2[10], "cong ty cp nguoi ban so thanh toan luong thang 10 2026")

        # Row 3: Huynh Tran Quang Khanh (ACB Bank)
        r3 = rows[2]
        self.assertEqual(r3[0], 3, "STT must be 3")
        self.assertEqual(r3[2], "12154887")
        self.assertEqual(r3[6], "HUYNH TRAN QUANG KHANH")
        self.assertEqual(r3[7], "ACB", "Bank must be identified as ACB")
        self.assertEqual(str(r3[8]), "24876000")

        # Test single payslip export
        slip_bytes, slip_fname = self.slip_hoang.generate_vcb_digibiz_xlsx()
        self.assertTrue(slip_bytes)
        wb_single = openpyxl.load_workbook(io.BytesIO(slip_bytes))
        ws_single = wb_single["Sheet1"]
        self.assertEqual(ws_single.cell(row=2, column=7).value, "LE XUAN HOANG")

        # Test action returning act_url
        act = self.pay_run.action_export_vcb_digibiz()
        self.assertEqual(act['type'], 'ir.actions.act_url')
        self.assertIn('/web/content/', act['url'])

    # =========================================================================
    # 2. LEGAL ENTITY & ADDRESS COMPLIANCE SCANNER TESTS
    # =========================================================================
    def test_02_legal_entity_and_address_compliance_scanner(self):
        """Verify the scanner detects and flags:
        - Legacy entity name (0313683499 / Ứng Dụng Sáng Tạo)
        - Outdated ward address (Phường Tân Phú, Quận 7 instead of Phường Tân Mỹ)
        - Allows clean documents conforming to Digimate JSC (MST: 0317958465, Phường Tân Mỹ).
        """
        # Test Case A: Legacy Entity Name
        text_legacy_name = "Hợp đồng dịch vụ ký kết bởi Công ty TNHH Ứng Dụng Sáng Tạo ngày 15/01/2023."
        res_a = self.scanner.scan_text(text_legacy_name)
        self.assertFalse(res_a['is_compliant'])
        self.assertTrue(res_a['has_legacy_entity'])
        self.assertIn("Ứng Dụng Sáng Tạo", res_a['summary'])

        # Test Case A2: Inverted Legacy Entity Name Variants (e.g. Sáng Tạo Ứng Dụng)
        for inverted_name in ("Sáng Tạo Ứng Dụng", "Sang Tao Ung Dung", "Công ty TNHH Sáng Tạo Ứng Dụng"):
            text_inverted = f"Đơn vị phát hành: {inverted_name} hợp đồng số 42/2023."
            res_inv = self.scanner.scan_text(text_inverted)
            self.assertFalse(res_inv['is_compliant'], f"Inverted entity name '{inverted_name}' must be flagged non-compliant")
            self.assertTrue(res_inv['has_legacy_entity'], f"Inverted entity name '{inverted_name}' must flag has_legacy_entity")

        # Test Case B: Legacy Tax Code
        text_legacy_mst = "Bên A: Đại diện công ty mã số thuế 0313683499 tại TP.HCM."
        res_b = self.scanner.scan_text(text_legacy_mst)
        self.assertFalse(res_b['is_compliant'])
        self.assertTrue(res_b['has_legacy_tax_code'])
        self.assertIn("0313683499", res_b['summary'])

        # Test Case B2: Formatted Legacy Tax Code Variants (hyphens, dots, spaces)
        for formatted_vat in ("0313-683-499", "0313.683.499", "0313 683 499"):
            text_vat = f"Bên A: Mã số thuế cũ {formatted_vat} cần rà soát chuyển đổi."
            res_vat = self.scanner.scan_text(text_vat)
            self.assertFalse(res_vat['is_compliant'], f"Formatted tax code '{formatted_vat}' must be flagged non-compliant")
            self.assertTrue(res_vat['has_legacy_tax_code'], f"Formatted tax code '{formatted_vat}' must flag has_legacy_tax_code")
            self.assertIn(formatted_vat, res_vat['summary'])

        # Test Case C: Outdated Ward Address
        text_outdated_ward = "Trụ sở tại: Số 11, Đường số 3, KDC CityLand, Phường Tân Phú, Quận 7, TP. Hồ Chí Minh."
        res_c = self.scanner.scan_text(text_outdated_ward)
        self.assertFalse(res_c['is_compliant'])
        self.assertTrue(res_c['has_outdated_ward'])
        self.assertIn("Phường Tân Mỹ", res_c['summary'])

        # Test Case D: Fully Compliant Document
        text_compliant = (
            "CÔNG TY CỔ PHẦN NGƯỜI BẠN SỐ (DIGIMATE JSC)\n"
            "Mã số thuế: 0317958465\n"
            "Địa chỉ: Số 11, Đường số 3, KDC CityLand, Phường Tân Mỹ, TP. Hồ Chí Minh."
        )
        res_d = self.scanner.scan_text(text_compliant)
        self.assertTrue(res_d['is_compliant'])
        self.assertFalse(res_d['has_legacy_entity'])
        self.assertFalse(res_d['has_legacy_tax_code'])
        self.assertFalse(res_d['has_outdated_ward'])
        self.assertEqual(len(res_d['violations']), 0)

        # Test Case E: Document model audit
        doc_deviated = self.doc_model.create({
            'name': 'HĐLĐ_Cu_Ung_Dung_Sang_Tao.txt',
            'raw': b"Cong ty TNHH Ung Dung Sang Tao cam ket chi tra...",
            'mimetype': 'text/plain',
        })
        doc_deviated.action_scan_single_entity_compliance()
        self.assertEqual(doc_deviated.al2c_compliance_status, 'deviated')
        self.assertTrue(doc_deviated.al2c_compliance_summary)

        # Test Case F: Sign Request audit
        sign_template = self.env['sign.template'].search([], limit=1)
        if not sign_template:
            sign_template = self.env['sign.template'].create({'name': 'Test NDA Template'})
        sign_req = self.env['sign.request'].create({
            'template_id': sign_template.id,
            'reference': 'Thỏa thuận bảo mật Ứng Dụng Sáng Tạo',
            'message': 'Kính gửi chi cục thuế mã số 0313683499 tại Phường Tân Phú, Quận 7',
        })
        req_res = sign_req.action_scan_single_entity_compliance()
        self.assertFalse(req_res['is_compliant'])
        self.assertTrue(sign_req.al2c_has_deviation)
        self.assertEqual(sign_req.al2c_clause_status, 'deviated')

    # =========================================================================
    # 3. FOREIGN REMITTANCE 24H SLA INVOICE MONITOR TESTS
    # =========================================================================
    def test_03_foreign_export_invoice_24h_sla_scanner(self):
        """Verify the compliance scanner monitors foreign remittances (Canva USD wires)
        and flags SLA breaches (> 24h past bank credit date without issued e-invoice).
        """
        usd = self.env.ref('base.USD', raise_if_not_found=False) or self.env['res.currency'].search([('name', '=', 'USD')], limit=1)
        now = fields.Datetime.now()
        past_36h = now - datetime.timedelta(hours=36)
        recent_2h = now - datetime.timedelta(hours=2)

        # 1. Canva Wire Overdue (36h past credit, no e-invoice) -> MUST FLAG
        move_breached = self.env['account.move'].create({
            'name': 'Canva USD Payout Award - Oct 2026',
            'move_type': 'out_invoice',
            'currency_id': usd.id if usd else self.company.currency_id.id,
            'al2c_is_foreign_remittance': True,
            'al2c_bank_credit_date': past_36h,
            'al2c_einvoice_number': False,
            'al2c_einvoice_state': 'draft',
            'company_id': self.company.id,
        })

        # 2. Canva Wire Recent (2h past credit, within 24h SLA) -> MUST NOT FLAG
        move_recent = self.env['account.move'].create({
            'name': 'Canva USD Wire Payout - Fresh Receipt',
            'move_type': 'out_invoice',
            'currency_id': usd.id if usd else self.company.currency_id.id,
            'al2c_is_foreign_remittance': True,
            'al2c_bank_credit_date': recent_2h,
            'al2c_einvoice_number': False,
            'al2c_einvoice_state': 'draft',
            'company_id': self.company.id,
        })

        # 3. Canva Wire Overdue BUT E-Invoice Already Issued -> MUST NOT FLAG
        move_issued = self.env['account.move'].create({
            'name': 'Canva USD Wire - Issued Statutory E-Invoice',
            'move_type': 'out_invoice',
            'currency_id': usd.id if usd else self.company.currency_id.id,
            'al2c_is_foreign_remittance': True,
            'al2c_bank_credit_date': past_36h,
            'al2c_einvoice_number': '1C26TAA-00000042',
            'al2c_einvoice_state': 'issued',
            'company_id': self.company.id,
        })

        # Run scanner
        res = self.scanner.scan_export_invoices_sla(hours_threshold=24)
        self.assertGreaterEqual(res['breached_count'], 1)
        self.assertIn(move_breached.id, res['breached_moves'].ids)
        self.assertNotIn(move_recent.id, res['breached_moves'].ids)
        self.assertNotIn(move_issued.id, res['breached_moves'].ids)

        self.assertTrue(move_breached.al2c_sla_breached, "Overdue foreign remittance must be flagged")
        self.assertFalse(move_recent.al2c_sla_breached, "Recent foreign remittance must not be flagged")
        self.assertFalse(move_issued.al2c_sla_breached, "Issued remittance must not be flagged")

    # =========================================================================
    # 4. DGM FOLDER WORKSPACE RULES & AI SORT PROMPT TESTS
    # =========================================================================
    def test_04_dgm_workspace_ai_sort_prompt_and_automation(self):
        """Verify configuration of ai_sort_prompt and folder automation rules on DGM
        documents.document workspace folders (3223 and children).
        """
        cfg = self.doc_model.configure_dgm_workspace_ai_rules()
        self.assertTrue(cfg, "Configuration routine must execute successfully")

        root = self.doc_model.browse(3223)
        if not root.exists() or root.type != 'folder':
            root = self.doc_model.search([('name', 'ilike', 'DGM · Công ty Cổ phần Người Bạn Số'), ('type', '=', 'folder')], limit=1)

        self.assertTrue(root.exists(), "Root folder must exist")
        self.assertTrue(root.ai_sort_prompt, "Root folder must have ai_sort_prompt configured")

        # Verify presence of key destination tags in prompt
        prompt = root.ai_sort_prompt
        self.assertIn("data-ai-record-id", prompt, "Prompt must embed valid target folder tags")
        self.assertIn("Canva", prompt)
        self.assertIn("BHXH", prompt)
        self.assertIn("DigiBiz", prompt)
        self.assertIn("Phường Tân Mỹ", prompt)

        # Test document routing logic
        doc_canva = self.doc_model.create({
            'name': 'Canva_USD_Wire_Payout_Notice_Oct2026.pdf',
            'mimetype': 'application/pdf',
        })
        doc_canva.classify_dgm_document()
        if self.doc_model.browse(3249).exists():
            self.assertEqual(doc_canva.folder_id.id, 3249, "Canva voucher must route to Hoá đơn đầu ra (3249)")

        doc_bhxh = self.doc_model.create({
            'name': 'Thong_bao_C12_TS_BHXH_TZ0857G_T10_2026.pdf',
            'mimetype': 'application/pdf',
        })
        doc_bhxh.classify_dgm_document()
        if self.doc_model.browse(3248).exists():
            self.assertEqual(doc_bhxh.folder_id.id, 3248, "C12-TS notice must route to BHXH (3248)")

        doc_vcb = self.doc_model.create({
            'name': 'VCB_DigiBiz_Bang_Luong_Thang_10_2026.xlsx',
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        doc_vcb.classify_dgm_document()
        if self.doc_model.browse(3252).exists():
            self.assertEqual(doc_vcb.folder_id.id, 3252, "VCB DigiBiz sheet must route to Lương DGM (3252)")

        doc_vnpt = self.doc_model.create({
            'name': 'Hoa_don_dau_vao_VNPT_Internet_T10_2026.pdf',
            'mimetype': 'application/pdf',
        })
        doc_vnpt.classify_dgm_document()
        if self.doc_model.browse(3250).exists():
            self.assertEqual(doc_vnpt.folder_id.id, 3250, "VNPT bill must route to Hoá đơn đầu vào (3250)")

        doc_contract = self.doc_model.create({
            'name': 'Hop_dong_lao_dong_2026_Le_Xuan_Hoang.docx',
            'mimetype': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        })
        doc_contract.classify_dgm_document()
        if self.doc_model.browse(3227).exists():
            self.assertEqual(doc_contract.folder_id.id, 3227, "Labor contract must route to Nhân sự (3227)")

    # =========================================================================
    # 5. SOP-RAG INDEXING & RETRIEVAL VALIDATION TESTS
    # =========================================================================
    def test_05_sop_rag_indexing_and_retrieval(self):
        """Verify reindex_dgm_sops() indexes DGM knowledge articles into al2c.rag.chunk
        with citation tags [Doc | Sec | ID | Hash] and validates retrieval for:
        - Canva 0% export invoicing
        - MTI overtime tiers
        - VCB DigiBiz 11 columns
        - BHXH TZ0857G
        - Mojo AI gross-to-net accounting
        """
        # Execute DGM SOP re-indexing
        report = self.rag_service.reindex_dgm_sops()
        self.assertGreaterEqual(report['articles_indexed'], 1, "Must index at least 1 DGM SOP article")
        self.assertGreaterEqual(report['chunks_created'], 1, "Must generate valid RAG chunks")

        # Verify citation tag format on stored chunks
        Chunk = self.env['al2c.rag.chunk'].sudo()
        dgm_chunks = Chunk.search([('source_model', '=', 'knowledge.article')], limit=20)
        self.assertTrue(dgm_chunks, "Chunks must be persisted in database")

        citation_regex = re.compile(r'^\[.* \| .* \| \d+ \| [0-9a-f]{8}\]$')
        for chunk in dgm_chunks:
            self.assertTrue(chunk.citation_tag, "Chunk citation tag must not be empty")
            self.assertTrue(
                citation_regex.match(chunk.citation_tag),
                f"Citation '{chunk.citation_tag}' must conform to [Doc | Sec | ID | Hash]"
            )

        # Validate retrieval across all 5 required operational domains
        # 1. Canva 0% export invoicing
        res_canva = self.rag_service.retrieve_grounded_context("Canva 0% export invoicing", top_k=3)
        self.assertTrue(len(res_canva) >= 1, "Must retrieve chunk for Canva 0% export invoicing")
        canva_text = " ".join(res_canva.mapped('content')).lower()
        self.assertTrue(any(term in canva_text for term in ('canva', 'xuất khẩu', '0%')), "Content must discuss Canva export")

        # 2. MTI overtime tiers
        res_mti = self.rag_service.retrieve_grounded_context("MTI overtime tiers", top_k=3)
        self.assertTrue(len(res_mti) >= 1, "Must retrieve chunk for MTI overtime tiers")
        mti_text = " ".join(res_mti.mapped('content')).lower()
        self.assertTrue(any(term in mti_text for term in ('mti', 'overtime', '100%', '150%')), "Content must discuss MTI OT tiers")

        # 3. VCB DigiBiz 11 columns
        res_vcb = self.rag_service.retrieve_grounded_context("VCB DigiBiz 11 columns", top_k=3)
        self.assertTrue(len(res_vcb) >= 1, "Must retrieve chunk for VCB DigiBiz 11 columns")
        vcb_text = " ".join(res_vcb.mapped('content')).lower()
        self.assertTrue(any(term in vcb_text for term in ('digibiz', 'vcb', 'vietcombank', '11')), "Content must discuss VCB DigiBiz")

        # 4. BHXH TZ0857G
        res_bhxh = self.rag_service.retrieve_grounded_context("BHXH TZ0857G", top_k=3)
        self.assertTrue(len(res_bhxh) >= 1, "Must retrieve chunk for BHXH TZ0857G")
        bhxh_text = " ".join(res_bhxh.mapped('content')).lower()
        self.assertTrue(any(term in bhxh_text for term in ('tz0857g', 'bhxh', 'c12-ts', 'quận 7')), "Content must discuss BHXH TZ0857G")

        # 5. Mojo AI gross-to-net accounting
        res_mojo = self.rag_service.retrieve_grounded_context("Mojo AI gross-to-net accounting", top_k=3)
        self.assertTrue(len(res_mojo) >= 1, "Must retrieve chunk for Mojo AI gross-to-net accounting")
        mojo_text = " ".join(res_mojo.mapped('content')).lower()
        self.assertTrue(any(term in mojo_text for term in ('mojo', 'gross', '511', 'app store')), "Content must discuss Mojo AI accounting")

    # =========================================================================
    # 6. ZERO FORBIDDEN BRANDING AUDIT
    # =========================================================================
    def test_06_zero_forbidden_branding(self):
        """Mandatory constraint: ZERO occurrences of forbidden legacy platform strings
        across user-facing outputs and templates.
        """
        # Check generated VCB Excel narration
        excel_bytes, _ = self.pay_run.generate_vcb_digibiz_xlsx()
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb["Sheet1"]
        for row in ws.iter_rows(values_only=True):
            for cell_val in row:
                if cell_val and isinstance(cell_val, str):
                    self.assertNotIn("odoo", cell_val.lower(), "Zero forbidden branding in VCB export")
                    self.assertNotIn("openerp", cell_val.lower(), "Zero forbidden branding in VCB export")

        # Check scanner summary messages
        res = self.scanner.scan_text("Văn bản hành chính hợp chuẩn DGM 0317958465.")
        self.assertNotIn("odoo", res['summary'].lower())
        self.assertNotIn("openerp", res['summary'].lower())

    # =========================================================================
    # 7. MULTIPAGE PDF 'TO SPLIT' TAG & CLASSIFICATION TESTS
    # =========================================================================
    def test_07_multipage_pdf_to_split_tag_assignment(self):
        """Verify that multipage PDF documents (is_multipage=True) automatically receive
        the 'To Split' tag (auto-created if missing) and remain unrouted in folder None.
        """
        # Build genuine 2-page PDF
        writer = PdfFileWriter()
        writer.addBlankPage(width=200, height=200)
        writer.addBlankPage(width=200, height=200)
        buf = io.BytesIO()
        writer.write(buf)
        multipage_pdf = buf.getvalue()

        # 1. Existing tag scenario
        tag_to_split = self.env['documents.tag'].search([('name', 'ilike', 'To Split')], limit=1)
        if not tag_to_split:
            tag_to_split = self.env['documents.tag'].create({'name': 'To Split'})

        doc_multi = self.doc_model.create({
            'name': 'Canva_MultiPage_Remittance_Batch.pdf',
            'mimetype': 'application/pdf',
            'raw': multipage_pdf,
        })
        self.assertTrue(doc_multi._get_is_multipage(), "Must identify genuine 2-page PDF as multipage")

        # Classify document
        doc_multi.classify_dgm_document()

        # Verify tag 'To Split' is attached and document stays in folder None
        self.assertIn(tag_to_split, doc_multi.tag_ids, "Multipage doc must receive 'To Split' tag")
        self.assertFalse(doc_multi.folder_id, "Multipage doc must remain unrouted in folder None")

        # 2. Auto-creation scenario when 'To Split' tag is missing
        all_to_split = self.env['documents.tag'].search([('name', 'ilike', 'To Split')])
        all_to_split.unlink()
        self.assertFalse(self.env['documents.tag'].search([('name', 'ilike', 'To Split')]))

        doc_multi_auto = self.doc_model.create({
            'name': 'VNPT_MultiPage_Vendor_Bills.pdf',
            'mimetype': 'application/pdf',
            'raw': multipage_pdf,
        })
        self.assertTrue(doc_multi_auto._get_is_multipage(), "Must identify genuine 2-page PDF as multipage")
        doc_multi_auto.classify_dgm_document()

        created_tag = self.env['documents.tag'].search([('name', 'ilike', 'To Split')], limit=1)
        self.assertTrue(created_tag.exists(), "Tag 'To Split' must be auto-created on demand")
        self.assertEqual(created_tag.name, 'To Split')
        self.assertIn(created_tag, doc_multi_auto.tag_ids, "Document must have auto-created tag attached")
        self.assertFalse(doc_multi_auto.folder_id, "Multipage doc must remain unrouted in folder None")

        # 3. Vectorized batch classification scenario on multiple documents simultaneously
        doc_batch_1 = self.doc_model.create({
            'name': 'Batch_MultiPage_1.pdf',
            'mimetype': 'application/pdf',
            'raw': multipage_pdf,
        })
        doc_batch_2 = self.doc_model.create({
            'name': 'Batch_MultiPage_2.pdf',
            'mimetype': 'application/pdf',
            'raw': multipage_pdf,
        })
        batch_docs = doc_batch_1 | doc_batch_2
        batch_docs.classify_dgm_document()
        for doc in batch_docs:
            self.assertIn(created_tag, doc.tag_ids, "All batch documents must receive 'To Split' tag in vectorized pass")
            self.assertFalse(doc.folder_id, "All batch documents must remain unrouted")

    def test_08_dgm_autonomous_crons_registration_and_execution(self):
        """Verify the 3 DGM autonomous watchdog crons:
        1. ir_cron_dgm_document_triage_split
        2. ir_cron_dgm_legal_entity_sla_watchdog
        3. ir_cron_dgm_bhxh_tk242_watchdog
        Assert proper model binding, active status, zero branding, clean execution,
        and kill-switch behavior when autonomy is disabled.
        """
        ICP = self.env['ir.config_parameter'].sudo()

        # 1. Assert al2c.autonomy_enabled and dgm.autonomy_enabled exist and are True
        al2c_on = ICP.get_param('al2c.autonomy_enabled')
        self.assertEqual(al2c_on, 'True', "al2c.autonomy_enabled must be True")
        dgm_on = ICP.get_param('dgm.autonomy_enabled')
        self.assertEqual(dgm_on, 'True', "dgm.autonomy_enabled must be True")
        self.assertTrue(self.env['al2c.autonomy.guard']._dgm_autonomy_enabled(), "Guard helper _dgm_autonomy_enabled must be True")

        # 2. Check Cron 1: Document Triage & Multipage Split Watchdog
        cron_triage = self.env.ref('insilos_al2c.ir_cron_dgm_document_triage_split', raise_if_not_found=False)
        self.assertTrue(cron_triage, "Cron ir_cron_dgm_document_triage_split must exist")
        self.assertTrue(cron_triage.active, "Document triage cron must be active")
        self.assertEqual(cron_triage.model_id.model, 'documents.document')
        self.assertNotIn("odoo", cron_triage.name.lower(), "Zero branding: 'odoo' forbidden in cron name")
        self.assertNotIn("openerp", cron_triage.name.lower(), "Zero branding: 'openerp' forbidden in cron name")
        res_triage = self.env['documents.document'].cron_dgm_document_triage_split()
        self.assertEqual(res_triage.get('status'), 'success', "Document triage cron must return success")
        self.assertIn('scanned', res_triage)
        self.assertIn('routed', res_triage)
        self.assertIn('multipage_split', res_triage)

        # 3. Check Cron 2: Legal Entity & 24h SLA Watchdog
        cron_legal = self.env.ref('insilos_al2c.ir_cron_dgm_legal_entity_sla_watchdog', raise_if_not_found=False)
        self.assertTrue(cron_legal, "Cron ir_cron_dgm_legal_entity_sla_watchdog must exist")
        self.assertTrue(cron_legal.active, "Legal entity SLA cron must be active")
        self.assertEqual(cron_legal.model_id.model, 'al2c.compliance.scanner')
        self.assertNotIn("odoo", cron_legal.name.lower(), "Zero branding: 'odoo' forbidden in cron name")
        self.assertNotIn("openerp", cron_legal.name.lower(), "Zero branding: 'openerp' forbidden in cron name")
        res_legal = self.env['al2c.compliance.scanner'].cron_dgm_legal_entity_sla_watchdog()
        self.assertEqual(res_legal.get('status'), 'success', "Legal entity SLA watchdog must return success")
        self.assertIn('documents_scanned', res_legal)
        self.assertIn('documents_deviated', res_legal)
        self.assertIn('contracts_scanned', res_legal)
        self.assertIn('export_sla_breaches', res_legal)

        # 4. Check Cron 3: BHXH & TK 242 Watchdog
        cron_bhxh = self.env.ref('insilos_al2c.ir_cron_dgm_bhxh_tk242_watchdog', raise_if_not_found=False)
        self.assertTrue(cron_bhxh, "Cron ir_cron_dgm_bhxh_tk242_watchdog must exist")
        self.assertTrue(cron_bhxh.active, "BHXH & TK 242 cron must be active")
        self.assertEqual(cron_bhxh.model_id.model, 'al2c.compliance.scanner')
        self.assertNotIn("odoo", cron_bhxh.name.lower(), "Zero branding: 'odoo' forbidden in cron name")
        self.assertNotIn("openerp", cron_bhxh.name.lower(), "Zero branding: 'openerp' forbidden in cron name")
        res_bhxh = self.env['al2c.compliance.scanner'].cron_dgm_bhxh_tk242_watchdog()
        self.assertEqual(res_bhxh.get('status'), 'success', "BHXH & TK 242 watchdog must return success")
        self.assertIn('bhxh_zero_arrears_verified', res_bhxh)
        self.assertIn('tk242_schedules_checked', res_bhxh)

        # 5. Check Kill-Switch Graceful Behavior
        ICP.set_param('dgm.autonomy_enabled', 'False')
        self.assertFalse(self.env['al2c.autonomy.guard']._dgm_autonomy_enabled())
        self.assertEqual(self.env['documents.document'].cron_dgm_document_triage_split().get('status'), 'skipped')
        self.assertEqual(self.env['al2c.compliance.scanner'].cron_dgm_legal_entity_sla_watchdog().get('status'), 'skipped')
        self.assertEqual(self.env['al2c.compliance.scanner'].cron_dgm_bhxh_tk242_watchdog().get('status'), 'skipped')

        # Restore kill-switch
        ICP.set_param('dgm.autonomy_enabled', 'True')
        self.assertTrue(self.env['al2c.autonomy.guard']._dgm_autonomy_enabled())

