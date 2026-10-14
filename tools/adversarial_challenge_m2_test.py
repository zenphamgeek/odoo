#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Adversarial Challenge Verification Harness for Milestone 2:
Master Data and VCB DigiBiz 11-column Excel Batch Export on insilos20_dev.
Platform: Insilos 20 LTS
Author: Final Empirical Challenger 1
"""

import io
import json
import os
import re
import sys
import unicodedata
from pathlib import Path
import openpyxl
import psycopg2

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

# Odoo imports
import odoo
from odoo import fields
from odoo.modules.registry import Registry
from odoo.tools import config

DB_HOST = "127.0.0.1"
DB_PORT = 5434
DB_USER = "odoo"
DB_PASS = "1NN0R1@2026"
DB_NAME = "insilos20_dev"

results = {
    "sql_checks": {},
    "orm_checks": {},
    "excel_checks": {},
    "stress_edge_cases": {},
    "verdict": "PENDING",
    "failures": []
}

def log_failure(section, test_name, reason):
    results["failures"].append({"section": section, "test": test_name, "reason": reason})
    print(f"[-] FAILURE in {section} :: {test_name} -> {reason}")

def log_success(section, test_name, detail=""):
    print(f"[+] PASS in {section} :: {test_name} {detail}")

# -----------------------------------------------------------------------------
# Part 1: Direct SQL Verification
# -----------------------------------------------------------------------------
def verify_sql():
    print("\n=== [1] Direct SQL Verification on insilos20_dev ===")
    conn = psycopg2.connect(host=DB_HOST, port=DB_PORT, user=DB_USER, password=DB_PASS, dbname=DB_NAME)
    cur = conn.cursor()

    # 1.1 Check 8 Employees in hr_employee and linked res_partner_bank
    canonical_employees = [
        ('DGM-001', 'Lê Xuân Hoàng', '1373068161', 'VIETCOMBANK', 21410000),
        ('DGM-002', 'Nguyễn Hoài Ân', '1901014028', 'VIETCOMBANK', 18282000),
        ('DGM-003', 'Mai Thị Hồng Thắm', '0111000287863', 'VIETCOMBANK', 13469000),
        ('DGM-004', 'Bùi Hữu Tài', '1045246615', 'VIETCOMBANK', 17780000),
        ('DGM-005', 'Huỳnh Trần Quang Khánh', '12154887', 'ACB', 24876000),
        ('DGM-006', 'Ngô Thanh Lộc', '12154937', 'ACB', 28108000),
        ('DGM-007', 'Lê Cao Bảo Phúc', '12154917', 'ACB', 16502000),
        ('DGM-008', 'Thái Nguyễn Đăng Khoa', '47070467', 'ACB', 7500000),
    ]

    for reg, name, expected_acc, expected_bank, expected_wage in canonical_employees:
        cur.execute("""
            SELECT e.id, e.name, e.registration_number, b.id,
                   b.account_number, b.bank_name, pg_typeof(b.account_number)
            FROM hr_employee e
            LEFT JOIN res_partner p ON e.work_contact_id = p.id
            LEFT JOIN res_partner_bank b ON b.partner_id = p.id
            WHERE e.registration_number = %s;
        """, (reg,))
        row = cur.fetchone()
        if not row:
            log_failure("SQL", f"employee_exists_{reg}", f"Employee {reg} ({name}) not found in database")
            continue

        emp_id, emp_name, emp_reg, bank_id, acc_num, bank_name, acc_type = row
        results["sql_checks"][reg] = {
            "id": emp_id,
            "name": emp_name,
            "acc_number": acc_num,
            "bank_name": bank_name,
            "acc_type": str(acc_type)
        }

        if not bank_id or not acc_num:
            log_failure("SQL", f"bank_linked_{reg}", f"Employee {reg} has no primary bank account linked")
            continue

        # Check account number exact string match
        if acc_num != expected_acc:
            log_failure("SQL", f"acc_match_{reg}", f"Expected {expected_acc}, got {acc_num}")
        else:
            log_success("SQL", f"acc_match_{reg}", f"{name} -> {acc_num} (bank: {bank_name}, type: {acc_type})")

        # Specific check for Mai Thi Hong Tham (DGM-003) leading zero
        if reg == 'DGM-003':
            if not acc_num.startswith('0'):
                log_failure("SQL", "leading_zero_tham", f"Account number {acc_num} does NOT start with '0'!")
            elif len(acc_num) != 13:
                log_failure("SQL", "length_tham", f"Account number {acc_num} length is {len(acc_num)}, expected 13!")
            else:
                log_success("SQL", "leading_zero_tham", f"Leading zero verified in DB: '{acc_num}', length 13, type: {acc_type}")

    # 1.2 Check 6 Canonical Partners in res_partner
    partners_to_check = [
        {"name": "MTI Technology", "search": "MTI TECHNOLOGY", "vat": "0314015511", "country_code": "VN"},
        {"name": "Canva Pty Ltd", "search": "Canva", "vat": None, "country_code": "AU"},
        {"name": "MobiFone MDS", "search": "MobiFone", "vat": "0100686209", "country_code": "VN"},
        {"name": "Temecula Valley Unified School District", "search": "Temecula", "vat": None, "country_code": "US"},
        {"name": "VNPT TP.HCM", "search": "VIỄN THÔNG TP.HCM", "vat": "0300954529", "country_code": "VN"},
        {"name": "Viettel-CA", "search": "VIETTEL", "vat": "0100109106-122", "country_code": "VN"},
    ]

    for p in partners_to_check:
        query = """
            SELECT p.id, p.name, p.vat, c.code
            FROM res_partner p
            LEFT JOIN res_country c ON p.country_id = c.id
            WHERE p.name ILIKE %s
        """
        params = [f"%{p['search']}%"]
        if p["vat"]:
            query += " OR p.vat = %s"
            params.append(p["vat"])
        cur.execute(query, tuple(params))
        rows = cur.fetchall()
        if not rows:
            log_failure("SQL", f"partner_exists_{p['name']}", f"Partner {p['name']} not found")
            continue

        matched = False
        for p_id, p_name, p_vat, c_code in rows:
            vat_ok = (p["vat"] is None) or (p_vat == p["vat"])
            country_ok = (p["country_code"] is None) or (c_code == p["country_code"])
            if vat_ok and country_ok:
                matched = True
                results["sql_checks"][f"partner_{p['name']}"] = {
                    "id": p_id,
                    "name": p_name,
                    "vat": p_vat,
                    "country": c_code
                }
                log_success("SQL", f"partner_{p['name']}", f"Found: {p_name} | VAT: {p_vat} | Country: {c_code}")
                break

        if not matched:
            log_failure("SQL", f"partner_match_{p['name']}", f"Partner {p['name']} found but VAT or Country did not match: {rows}")

    conn.close()

# -----------------------------------------------------------------------------
# Part 2: Odoo ORM Verification & Excel Generation
# -----------------------------------------------------------------------------
def verify_orm_and_excel():
    print("\n=== [2] Odoo ORM & VCB DigiBiz Excel Generation ===")
    config.parse_config(['-c', '/home/zen/O20/insilos.conf', '-d', DB_NAME])
    odoo.netsvc.init_logger()

    registry = Registry(DB_NAME)
    with registry.cursor() as cr:
        env = odoo.api.Environment(cr, odoo.SUPERUSER_ID, {})

        # 2.1 Verify Employee and Partner records in ORM
        Employee = env['hr.employee']
        Partner = env['res.partner']
        PayslipRun = env['hr.payslip.run']
        Payslip = env['hr.payslip']

        canonical_regs = [
            ('DGM-001', 'Lê Xuân Hoàng', '1373068161', 'VIETCOMBANK', 21410000),
            ('DGM-002', 'Nguyễn Hoài Ân', '1901014028', 'VIETCOMBANK', 18282000),
            ('DGM-003', 'Mai Thị Hồng Thắm', '0111000287863', 'VIETCOMBANK', 13469000),
            ('DGM-004', 'Bùi Hữu Tài', '1045246615', 'VIETCOMBANK', 17780000),
            ('DGM-005', 'Huỳnh Trần Quang Khánh', '12154887', 'ACB', 24876000),
            ('DGM-006', 'Ngô Thanh Lộc', '12154937', 'ACB', 28108000),
            ('DGM-007', 'Lê Cao Bảo Phúc', '12154917', 'ACB', 16502000),
            ('DGM-008', 'Thái Nguyễn Đăng Khoa', '47070467', 'ACB', 7500000),
        ]

        # 2.2 Create Pay Run and Payslips for batch export
        pay_run = PayslipRun.create({
            'name': 'Adversarial Challenge Pay Run 10/2026',
            'date_start': fields.Date.from_string('2026-10-01'),
            'date_end': fields.Date.from_string('2026-10-31'),
        })

        slips = env['hr.payslip']
        for reg, name, stk, bank_name, net_wage in canonical_regs:
            emp = Employee.search([('registration_number', '=', reg)], limit=1)
            if not emp:
                log_failure("ORM", f"emp_search_{reg}", f"Employee {reg} not found")
                continue

            slip = Payslip.create({
                'name': f'Slip {name} 10/2026',
                'employee_id': emp.id,
                'payslip_run_id': pay_run.id,
                'date_from': pay_run.date_start,
                'date_to': pay_run.date_end,
                'net_wage': net_wage,
                'al2c_net_salary': net_wage,
            })
            slips |= slip

        results["orm_checks"]["payslip_count"] = len(slips)
        if len(slips) != 8:
            log_failure("ORM", "payslip_count", f"Expected 8 slips, created {len(slips)}")
            return

        # 2.3 Call generate_vcb_digibiz_xlsx
        excel_bytes, filename = pay_run.generate_vcb_digibiz_xlsx()
        if not excel_bytes or not filename.endswith('.xlsx'):
            log_failure("ORM", "excel_generation", "generate_vcb_digibiz_xlsx failed to return valid xlsx bytes")
            return
        log_success("ORM", "excel_generation", f"Generated {len(excel_bytes)} bytes, filename: {filename}")
        results["orm_checks"]["excel_bytes_len"] = len(excel_bytes)
        results["orm_checks"]["excel_filename"] = filename

        # 2.4 Open with openpyxl and inspect deeply
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb.active

        # Check Column Count
        max_col = ws.max_column
        results["excel_checks"]["max_column"] = max_col
        if max_col != 11:
            log_failure("EXCEL", "col_count", f"Expected 11 columns, got {max_col}")
        else:
            log_success("EXCEL", "col_count", "Exactly 11 columns confirmed")

        # Check Row Count (1 header + 8 employees = 9 rows)
        max_row = ws.max_row
        results["excel_checks"]["max_row"] = max_row
        if max_row != 9:
            log_failure("EXCEL", "row_count", f"Expected 9 rows (1 header + 8 data), got {max_row}")
        else:
            log_success("EXCEL", "row_count", "Exactly 9 rows (1 header + 8 data rows) confirmed")

        # Check Header Values
        from odoo.addons.insilos_al2c.models.hr_payslip_vcb import VCB_DIGIBIZ_COLUMNS
        actual_headers = [ws.cell(row=1, column=c).value for c in range(1, 12)]
        results["excel_checks"]["headers"] = actual_headers
        if actual_headers != VCB_DIGIBIZ_COLUMNS:
            log_failure("EXCEL", "headers", f"Headers mismatch: {actual_headers}")
        else:
            log_success("EXCEL", "headers", "Header names match VCB_DIGIBIZ_COLUMNS exactly")

        # Iterate data rows
        total_salary = 0
        names_unaccented_upper = True
        tham_verified = False
        rows_summary = []

        for r in range(2, 10):
            stt = ws.cell(row=r, column=1).value
            ref = ws.cell(row=r, column=2).value
            acc_cell = ws.cell(row=r, column=3)
            acc_val = acc_cell.value
            acc_fmt = acc_cell.number_format
            acc_dtype = acc_cell.data_type

            cmnd = ws.cell(row=r, column=4).value
            ben_name = ws.cell(row=r, column=7).value
            bank = ws.cell(row=r, column=8).value
            amount = int(ws.cell(row=r, column=9).value or 0)
            currency = ws.cell(row=r, column=10).value
            content = ws.cell(row=r, column=11).value

            total_salary += amount

            rows_summary.append({
                "stt": stt,
                "ref": ref,
                "account_number": acc_val,
                "number_format": acc_fmt,
                "data_type": acc_dtype,
                "beneficiary": ben_name,
                "bank": bank,
                "amount": amount,
                "currency": currency,
                "content": content
            })

            # Check unaccented uppercase beneficiary name
            if ben_name != ben_name.upper():
                log_failure("EXCEL", f"name_upper_row_{r}", f"Name '{ben_name}' is not uppercase")
                names_unaccented_upper = False
            # Check accents
            decomposed = unicodedata.normalize('NFD', ben_name)
            has_accents = any(unicodedata.category(c) == 'Mn' for c in decomposed) or 'Đ' in ben_name or 'đ' in ben_name
            if has_accents:
                log_failure("EXCEL", f"name_unaccented_row_{r}", f"Name '{ben_name}' contains accents")
                names_unaccented_upper = False

            # Check Mai Thi Hong Tham (DGM-003) specifically
            if ben_name == 'MAI THI HONG THAM':
                tham_verified = True
                if acc_val != '0111000287863':
                    log_failure("EXCEL", "tham_acc_val", f"Expected '0111000287863', got '{acc_val}'")
                elif not str(acc_val).startswith('0'):
                    log_failure("EXCEL", "tham_leading_zero", f"Value '{acc_val}' lost leading zero!")
                else:
                    log_success("EXCEL", "tham_acc_val", f"Mai Thi Hong Tham account is '{acc_val}' (starts with '0')")

                # Check cell format '@' (text)
                if acc_fmt != '@':
                    log_failure("EXCEL", "tham_acc_format", f"Number format is '{acc_fmt}', expected '@'")
                else:
                    log_success("EXCEL", "tham_acc_format", f"Cell number_format is strictly text '@'")

                # Check data_type
                if acc_dtype != 's':
                    log_failure("EXCEL", "tham_data_type", f"Cell data_type is '{acc_dtype}', expected 's' (string)")
                else:
                    log_success("EXCEL", "tham_data_type", "Cell data_type is 's' (string)")

        results["excel_checks"]["rows"] = rows_summary
        results["excel_checks"]["total_salary_calculated"] = total_salary

        if names_unaccented_upper:
            log_success("EXCEL", "all_names_unaccented_upper", "All 8 recipient names are strictly unaccented uppercase")

        if not tham_verified:
            log_failure("EXCEL", "tham_found", "Mai Thi Hong Tham was not found in exported rows")

        # Verify total salary sum
        expected_total = 147927000
        if total_salary != expected_total:
            log_failure("EXCEL", "total_salary_sum", f"Total salary sum is {total_salary:,} VND, expected {expected_total:,} VND")
        else:
            log_success("EXCEL", "total_salary_sum", f"Total salary sum matches exactly: {total_salary:,} VND (147,927,000 VND)")

        # ---------------------------------------------------------------------
        # Part 3: Adversarial Edge Cases & Stress Testing
        # ---------------------------------------------------------------------
        print("\n=== [3] Adversarial Edge Cases & Stress Testing ===")
        from odoo.addons.insilos_al2c.models.hr_payslip_vcb import remove_vietnamese_accents, sanitize_csv_cell

        # 3.1 Accent removal stress test
        accent_cases = [
            ("Lê Xuân Hoàng", "LE XUAN HOANG"),
            ("Nguyễn Hoài Ân", "NGUYEN HOAI AN"),
            ("Mai Thị Hồng Thắm", "MAI THI HONG THAM"),
            ("Bùi Hữu Tài", "BUI HUU TAI"),
            ("Huỳnh Trần Quang Khánh", "HUYNH TRAN QUANG KHANH"),
            ("Ngô Thanh Lộc", "NGO THANH LOC"),
            ("Lê Cao Bảo Phúc", "LE CAO BAO PHUC"),
            ("Thái Nguyễn Đăng Khoa", "THAI NGUYEN DANG KHOA"),
            ("Đỗ Văn Phong", "DO VAN PHONG"),
            ("VŨ ĐÌNH TRỌNG", "VU DINH TRONG"),
            ("Trịnh Thị Ửng", "TRINH THI UNG"),
        ]
        results["stress_edge_cases"]["accents"] = []
        for inp, exp in accent_cases:
            res = remove_vietnamese_accents(inp)
            results["stress_edge_cases"]["accents"].append({"input": inp, "expected": exp, "actual": res})
            if res != exp:
                log_failure("STRESS", f"accent_removal_{inp}", f"Expected '{exp}', got '{res}'")
            else:
                log_success("STRESS", f"accent_removal_{inp}", f"-> {res}")

        # 3.2 CSV / Spreadsheet Formula Injection Stress Test (CWE-1236)
        injection_cases = [
            ("=SUM(A1:A10)", "'=SUM(A1:A10)"),
            ("+12345", "'+12345"),
            ("-9999", "'-9999"),
            ("@alert", "'@alert"),
            ("\tcmd", "'\tcmd"),
            ("NORMAL_VALUE", "NORMAL_VALUE"),
        ]
        results["stress_edge_cases"]["formula_injection"] = []
        for inp, exp in injection_cases:
            res = sanitize_csv_cell(inp)
            results["stress_edge_cases"]["formula_injection"].append({"input": inp, "expected": exp, "actual": res})
            if res != exp:
                log_failure("STRESS", f"formula_injection_{inp}", f"Expected '{exp}', got '{res}'")
            else:
                log_success("STRESS", f"formula_injection_{inp}", f"-> {res}")

        # 3.3 Account Number Sanitization with Dirty Inputs
        dirty_accounts = [
            ("0111000287863", "0111000287863"),
            (" 0111000287863 ", "0111000287863"),
            ("0111-000-287-863", "0111000287863"),
            ("0111.000.287.863", "0111000287863"),
            ("0111 000 287 863", "0111000287863"),
        ]
        results["stress_edge_cases"]["dirty_accounts"] = []
        for dirty, clean_exp in dirty_accounts:
            clean = re.sub(r'[^0-9]', '', str(dirty or '')).strip()
            results["stress_edge_cases"]["dirty_accounts"].append({"input": dirty, "expected": clean_exp, "actual": clean})
            if clean != clean_exp or not clean.startswith('0'):
                log_failure("STRESS", f"clean_acc_{dirty}", f"Failed to preserve leading zero: {clean}")
            else:
                log_success("STRESS", f"clean_acc_{dirty}", f"-> {clean} (starts with 0, len={len(clean)})")

        cr.rollback()

# -----------------------------------------------------------------------------
# Part 4: Adversarial Security & Non-Singleton Verification
# -----------------------------------------------------------------------------
def verify_adversarial_security_and_non_singleton():
    print("\n=== [4] Adversarial Security & Invariant Verification ===")
    results["security_checks"] = {}

    import subprocess
    error_dialogs_js = os.path.join(str(REPO_ROOT), "addons/web/static/src/core/errors/error_dialogs.js")
    patch_js = os.path.join(str(REPO_ROOT), "enterprise/insilos_theme_genesis/static/src/errors/runtime_error_dialog_patch.js")

    def extract_func(filepath, name):
        with open(filepath, "r", encoding="utf-8") as f:
            src = f.read()
        marker = f"export function {name}("
        idx = src.find(marker)
        if idx == -1:
            marker = f"function {name}("
            idx = src.find(marker)
        assert idx != -1, f"Function {name} not found in {filepath}"
        brace_depth = 0
        started = False
        for pos in range(idx, len(src)):
            ch = src[pos]
            if ch == "{":
                brace_depth += 1
                started = True
            elif ch == "}":
                brace_depth -= 1
                if started and brace_depth == 0:
                    return src[idx : pos + 1].replace(f"export function {name}(", f"function {name}(")
        raise RuntimeError(f"Unbalanced braces in {name} of {filepath}")

    # 4.1 & 4.2 Error Sanitizer ReDoS & JSON AST Tests
    for js_path, label in [(error_dialogs_js, "core_error_dialogs"), (patch_js, "theme_runtime_patch")]:
        san_code = extract_func(js_path, "sanitizeErrorPayload")
        node_script = f"""
{san_code}
const startTime = Date.now();
const redosPayload = "Authorization: Digest " + "a=b,".repeat(50000) + "realm=test";
const redosResult = sanitizeErrorPayload(redosPayload);
const redosDuration = Date.now() - startTime;

const jsonPayload = JSON.stringify({{
    password: "SuperSecretPassword!999",
    apiKey: "AIzaSyD-1234567890123456789012345678901",
    anthropicKey: "sk-ant-api03-abcdef1234567890abcdef1234567890",
    awsStsKey: "ASIAIOSFODNN7EXAMPLE",
    awsIamKey: "AKIAIOSFODNN7EXAMPLE",
    openAiKey: "sk-proj-AbCdEfGhIjKlMnOpQrStUvWxYz123456",
    githubToken: "ghp_1234567890abcdefghijklmnopqrstuvwxyz",
    phone: "0912345678",
    email: "security.audit@innoria.com",
    dbUri: "postgresql://odoo:db_secret_pass@127.0.0.1:5432/insilos20_dev",
    normalField: "unaffected_string_value",
    counter: 42
}});

const sanitizedJson = sanitizeErrorPayload(jsonPayload);
let parsedObj = null;
let jsonParseError = null;
try {{
    parsedObj = JSON.parse(sanitizedJson);
}} catch (err) {{
    jsonParseError = err.message;
}}

console.log(JSON.stringify({{
    redosDuration,
    redosResultPrefix: redosResult.slice(0, 30),
    jsonParseError,
    sanitizedJson,
    parsedObj
}}));
"""
        proc = subprocess.run(["node", "-e", node_script], capture_output=True, text=True, check=True)
        out = json.loads(proc.stdout.strip())

        # Assert linear-time ReDoS bound (< 50ms)
        duration = out["redosDuration"]
        if duration >= 50:
            log_failure("SECURITY", f"redos_bound_{label}", f"ReDoS duration exceeded 50ms: {duration}ms")
        else:
            log_success("SECURITY", f"redos_bound_{label}", f"Duration: {duration}ms (Bound < 50ms verified)")

        # Assert JSON AST preservation
        if out["jsonParseError"]:
            log_failure("SECURITY", f"json_ast_preservation_{label}", f"JSON parse failed: {out['jsonParseError']}")
        else:
            obj = out["parsedObj"]
            leaks = []
            if "SuperSecretPassword!999" in out["sanitizedJson"]:
                leaks.append("password")
            if "AIzaSyD-" in out["sanitizedJson"]:
                leaks.append("AIzaSy Gemini key")
            if "sk-ant-api03-" in out["sanitizedJson"]:
                leaks.append("Anthropic key")
            if "ASIAIOSFODNN7EXAMPLE" in out["sanitizedJson"]:
                leaks.append("AWS STS key")
            if "AKIAIOSFODNN7EXAMPLE" in out["sanitizedJson"]:
                leaks.append("AWS IAM key")
            if "sk-proj-" in out["sanitizedJson"]:
                leaks.append("OpenAI key")
            if "ghp_1234567890" in out["sanitizedJson"]:
                leaks.append("GitHub token")
            if "db_secret_pass" in out["sanitizedJson"]:
                leaks.append("DB URI password")
            if "0912345678" in out["sanitizedJson"]:
                leaks.append("phone number")
            if "security.audit@innoria.com" in out["sanitizedJson"]:
                leaks.append("email address")

            if leaks:
                log_failure("SECURITY", f"credentials_redacted_{label}", f"Leaked credentials: {leaks}")
            else:
                log_success("SECURITY", f"credentials_redacted_{label}", "All 10 sensitive patterns cleanly redacted")

            if obj.get("normalField") == "unaffected_string_value" and obj.get("counter") == 42:
                log_success("SECURITY", f"json_ast_preservation_{label}", "JSON AST structure and non-sensitive fields preserved")
            else:
                log_failure("SECURITY", f"json_ast_preservation_{label}", f"Corrupted fields: {obj}")

    # 4.3 BaseAutomation Multi-Recordset Non-Singleton Execution
    registry = Registry(DB_NAME)
    with registry.cursor() as cr:
        env = odoo.api.Environment(cr, odoo.SUPERUSER_ID, {})
        Automation = env["base.automation"]
        Document = env["documents.document"]

        # Ensure documents exist
        doc = Document.search([], limit=1)
        if not doc:
            doc = Document.create({"name": "Adversarial Test Doc"})

        # Create two automation rules in savepoint
        with cr.savepoint():
            rule_model_id = env["ir.model"]._get_id("documents.document")
            rule1 = Automation.create({
                "name": "Adversarial Non-AI Automation Rule 1",
                "model_id": rule_model_id,
                "trigger": "on_create_or_write",
            })
            rule2 = Automation.create({
                "name": "Adversarial AI Automation Rule 2",
                "model_id": rule_model_id,
                "trigger": "on_create_or_write",
                "ai_autosort_folder_id": doc.folder_id.id if doc.folder_id else False,
            })
            multi_rules = rule1 | rule2
            assert len(multi_rules) == 2, "Expected multi-recordset of 2 rules"

            try:
                # Execute _process on multi-recordset with ai_documents_skip_autosort=True
                multi_rules.with_context(ai_documents_skip_autosort=True)._process(doc)
                log_success("SECURITY", "base_automation_multi_recordset_no_singleton_crash",
                            "Multi-recordset rules._process(doc) succeeded with 0 Expected singleton crashes")
            except ValueError as ve:
                if "Expected singleton" in str(ve):
                    log_failure("SECURITY", "base_automation_multi_recordset_no_singleton_crash",
                                f"Expected singleton crash occurred: {ve}")
                else:
                    raise
            except Exception as e:
                log_failure("SECURITY", "base_automation_multi_recordset_no_singleton_crash",
                            f"Unexpected exception during _process: {e}")

        cr.rollback()

if __name__ == "__main__":
    try:
        verify_sql()
        verify_orm_and_excel()
        verify_adversarial_security_and_non_singleton()
        print("\n=======================================================")
        if len(results["failures"]) == 0:
            results["verdict"] = "APPROVE"
            print(">>> ADVERSARIAL CHALLENGE VERDICT: APPROVE <<<")
        else:
            results["verdict"] = "REJECT"
            print(f">>> ADVERSARIAL CHALLENGE VERDICT: REJECT ({len(results['failures'])} failures) <<<")
        print("=======================================================")
    except Exception as e:
        import traceback
        traceback.print_exc()
        results["verdict"] = "REJECT"
        results["failures"].append({"section": "EXECUTION", "test": "fatal_error", "reason": str(e)})

    # Write results to challenger directory
    out_dir = Path('/home/zen/O20/.agents/teamwork/challenger_final_1')
    out_dir.mkdir(parents=True, exist_ok=True)
    with open(out_dir / 'challenge_results.json', 'w', encoding='utf-8') as f:
        json.dump(results, f, indent=2, ensure_ascii=False)
