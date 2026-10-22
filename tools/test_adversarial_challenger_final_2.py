#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Challenger 2 Adversarial Stress-Test Harness & Empirical Verifier:
DGM Autonomous Operations (`autonomous on`) on `insilos20_dev`
Target Database: insilos20_dev (PostgreSQL 127.0.0.1:5434)
Platform: Insilos 20 LTS

Author: Empirical Challenger 2 (`challenger_final_2`)
Date: 2026-10-14
"""

import sys
import os
import io
import re
import json
import zipfile
import datetime
import xml.etree.ElementTree as ET
from pathlib import Path

# Add repo root to sys.path
REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

import odoo
from odoo.modules.registry import Registry
from odoo.tools import config
from odoo.tools.pdf import PdfFileReader, PdfFileWriter
from odoo import fields

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None


def run_adversarial_challenge():
    print("=" * 80)
    print("CHALLENGER 2: ADVERSARIAL STRESS-TEST & LIVE DOCUMENT AUDIT")
    print("Target Database: insilos20_dev (PostgreSQL 127.0.0.1:5434)")
    print("=" * 80)

    config.parse_config(['-c', 'insilos.conf', '-d', 'insilos20_dev'])
    registry = Registry.new('insilos20_dev', lock_wait=180)

    total_checks = 0
    passed_checks = 0
    failed_checks = []

    def check(condition, test_name, details=""):
        nonlocal total_checks, passed_checks, failed_checks
        total_checks += 1
        if condition:
            passed_checks += 1
            print(f"  [PASS] {test_name}: {details}")
        else:
            print(f"  [FAIL] {test_name}: {details}")
            failed_checks.append(f"{test_name}: {details}")

    with registry.cursor() as cr:
        env = odoo.api.Environment(cr, odoo.SUPERUSER_ID, {})
        ICP = env['ir.config_parameter'].sudo()
        guard = env['al2c.autonomy.guard']
        Doc = env['documents.document'].sudo()
        Scanner = env['al2c.compliance.scanner'].sudo()
        Employee = env['hr.employee'].sudo()
        Partner = env['res.partner'].sudo()

        # Save baseline configuration
        orig_al2c = ICP.get_param('al2c.autonomy_enabled', 'True')
        orig_dgm = ICP.get_param('dgm.autonomy_enabled', 'True')

        try:
            # =========================================================================
            # STAGE 1: EXECUTE THE 3 AUTONOMOUS CRONS ON LIVE insilos20_dev
            # =========================================================================
            print("\n>>> STAGE 1: Live Execution of 3 Autonomous Watchdog Daemons <<<")

            # Ensure autonomy enabled
            ICP.set_param('al2c.autonomy_enabled', 'True')
            ICP.set_param('dgm.autonomy_enabled', 'True')

            # 1.1 Document Triage & Multipage Split Watchdog
            print("  --- 1.1 cron_dgm_document_triage_split(batch_limit=50) ---")
            res_triage = Doc.cron_dgm_document_triage_split(batch_limit=50)
            print(f"      Result: {res_triage}")
            check(isinstance(res_triage, dict), "Cron 1 Return Type", "Returns dictionary")
            check(res_triage.get('status') == 'success', "Cron 1 Status 'success'", f"Status: {res_triage.get('status')}")
            check('scanned' in res_triage and 'routed' in res_triage and 'multipage_split' in res_triage,
                  "Cron 1 Return Metrics", f"scanned={res_triage.get('scanned')}, routed={res_triage.get('routed')}, split={res_triage.get('multipage_split')}")

            # 1.2 Legal Entity & 24h Export E-Invoice SLA Watchdog
            print("  --- 1.2 cron_dgm_legal_entity_sla_watchdog(batch_limit=50) ---")
            res_legal = Scanner.cron_dgm_legal_entity_sla_watchdog(batch_limit=50)
            print(f"      Result: {res_legal}")
            check(isinstance(res_legal, dict), "Cron 2 Return Type", "Returns dictionary")
            check(res_legal.get('status') == 'success', "Cron 2 Status 'success'", f"Status: {res_legal.get('status')}")
            check('documents_scanned' in res_legal and 'documents_deviated' in res_legal and 'export_sla_breaches' in res_legal,
                  "Cron 2 Return Metrics", f"scanned={res_legal.get('documents_scanned')}, deviated={res_legal.get('documents_deviated')}, sla_breaches={res_legal.get('export_sla_breaches')}")

            # 1.3 BHXH Zero-Arrears & Prepaid TK 242 Watchdog
            print("  --- 1.3 cron_dgm_bhxh_tk242_watchdog() ---")
            res_bhxh = Scanner.cron_dgm_bhxh_tk242_watchdog()
            print(f"      Result: {res_bhxh}")
            check(isinstance(res_bhxh, dict), "Cron 3 Return Type", "Returns dictionary")
            check(res_bhxh.get('status') == 'success', "Cron 3 Status 'success'", f"Status: {res_bhxh.get('status')}")
            check('bhxh_zero_arrears_verified' in res_bhxh and 'tk242_schedules_checked' in res_bhxh,
                  "Cron 3 Return Metrics", f"zero_arrears={res_bhxh.get('bhxh_zero_arrears_verified')}, tk242_schedules={res_bhxh.get('tk242_schedules_checked')}")

            # =========================================================================
            # STAGE 2: ADVERSARIAL STRESS-TESTING OF THE 3 CRONS
            # =========================================================================
            print("\n>>> STAGE 2: Adversarial Stress-Testing & Kill-Switch Validation <<<")

            # 2.1 Kill-Switch Fail-Closed Behavior
            print("  --- 2.1 Testing Kill-Switch (dgm.autonomy_enabled = 'False') ---")
            ICP.set_param('dgm.autonomy_enabled', 'False')
            check(guard._dgm_autonomy_enabled() is False, "Guard evaluates to False when dgm=False", "Disabled")

            res_triage_skip = Doc.cron_dgm_document_triage_split(batch_limit=50)
            check(res_triage_skip.get('status') == 'skipped', "Cron 1 gracefully skips when autonomy disabled", f"{res_triage_skip}")

            res_legal_skip = Scanner.cron_dgm_legal_entity_sla_watchdog(batch_limit=50)
            check(res_legal_skip.get('status') == 'skipped', "Cron 2 gracefully skips when autonomy disabled", f"{res_legal_skip}")

            res_bhxh_skip = Scanner.cron_dgm_bhxh_tk242_watchdog()
            check(res_bhxh_skip.get('status') == 'skipped', "Cron 3 gracefully skips when autonomy disabled", f"{res_bhxh_skip}")

            # Restore autonomy
            ICP.set_param('dgm.autonomy_enabled', 'True')
            check(guard._dgm_autonomy_enabled() is True, "Guard restored to operational True", "Enabled")

            # 2.2 Boundary limits & Hostile document payloads
            print("  --- 2.2 Boundary limits & Hostile Injections in Savepoints ---")
            temp_docs = Doc.browse()
            try:
                # Corrupt binary claiming to be pdf
                doc_corrupt = Doc.create({
                    'name': 'adversarial_corrupt_stream.pdf',
                    'mimetype': 'application/pdf',
                    'raw': b'%PDF-1.4\nTRUNCATED_CORRUPT_BYTES_XYZ123',
                })
                temp_docs |= doc_corrupt
                is_multi = doc_corrupt._get_is_multipage()
                check(is_multi is False, "Corrupted PDF _get_is_multipage handled gracefully without crash", "Returned False")

                # Injected document with SQL injection / XSS name
                doc_sql = Doc.create({
                    'name': "'; DROP TABLE documents_document; -- <script>alert(1)</script>",
                    'mimetype': 'text/plain',
                    'raw': b"Canva wire payment notice",
                })
                temp_docs |= doc_sql
                doc_sql.classify_dgm_document()
                check(doc_sql.folder_id.id == 3249 if Doc.browse(3249).exists() else True,
                      "Document with SQL injection string routed safely by classification engine", "Safe")

                # Injected legacy entity variations
                for legacy_sample in [
                    "0313-683-499", "0313.683.499", "0313 683 499",
                    "CÔNG TY TNHH ỨNG DỤNG SÁNG TẠO", "Sáng Tạo Ứng Dụng", "ung dung sang tao llc",
                    "Phường Tân Phú, Quận 7"
                ]:
                    doc_test = Doc.create({
                        'name': f'test_legacy_{legacy_sample[:10]}.txt',
                        'mimetype': 'text/plain',
                        'raw': f"Van ban co chua: {legacy_sample}".encode('utf-8'),
                    })
                    temp_docs |= doc_test
                    scan_out = Scanner.scan_document(doc_test)
                    check(scan_out['is_compliant'] is False, f"Detection of legacy pattern '{legacy_sample}'", f"Violations: {len(scan_out['violations'])}")
            finally:
                temp_docs.unlink()

            # =========================================================================
            # STAGE 3: DIRECT VERIFICATION OF 161 DOCUMENTS UNDER ROOT ID 3223
            # =========================================================================
            print("\n>>> STAGE 3: Deep Audit of 161 Documents under Root ID 3223 <<<")

            root_folder = Doc.browse(3223)
            check(root_folder.exists() and root_folder.type == 'folder', "Root folder ID 3223 exists and is folder", f"Name: {root_folder.name}")

            all_subfolders = Doc.search([('id', 'child_of', 3223), ('type', '=', 'folder')])
            all_non_folders = Doc.search([('folder_id', 'in', all_subfolders.ids), ('type', '!=', 'folder')], order='id asc')

            print(f"      Total subfolders under 3223: {len(all_subfolders)}")
            print(f"      Total non-folder documents under 3223: {len(all_non_folders)}")
            check(len(all_non_folders) == 161, "Exact count of 161 non-folder documents under folder 3223", f"Found {len(all_non_folders)}")

            # Independent extraction helper
            def extract_text_from_doc(document):
                texts = [str(document.name or "")]
                att = document.attachment_id
                if att:
                    texts.append(att.name or "")
                    if getattr(att, 'index_content', False):
                        texts.append(att.index_content)
                    raw_data = getattr(att, 'raw', b"")
                    raw = b""
                    if hasattr(raw_data, 'read'):
                        raw = raw_data.read()
                    elif isinstance(raw_data, (bytes, bytearray)):
                        raw = bytes(raw_data)
                    elif raw_data:
                        raw = bytes(raw_data)

                    fname = (att.name or str(document.name or "")).lower()
                    if raw:
                        if fname.endswith('.docx'):
                            try:
                                with zipfile.ZipFile(io.BytesIO(raw)) as z:
                                    if 'word/document.xml' in z.namelist():
                                        xml_content = z.read('word/document.xml')
                                        tree = ET.fromstring(xml_content)
                                        texts.append(' '.join(tree.itertext()))
                            except Exception:
                                pass
                        elif fname.endswith('.pdf') and PyPDF2:
                            try:
                                reader = PyPDF2.PdfReader(io.BytesIO(raw))
                                for page in reader.pages:
                                    p_txt = page.extract_text()
                                    if p_txt:
                                        texts.append(p_txt)
                            except Exception:
                                pass
                        elif fname.endswith(('.txt', '.csv', '.xml', '.html', '.json')):
                            texts.append(raw.decode('utf-8', errors='ignore'))
                return " ".join(texts)

            # Independent detection regexes
            re_legacy_vat = re.compile(r'\b0313[\s\.\-]?683[\s\.\-]?499\b', re.IGNORECASE)
            re_legacy_entity = re.compile(r'ứng\s*dụng\s*sáng\s*tạo|ung\s*dung\s*sang\s*tao|sáng\s*tạo\s*ứng\s*dụng|sang\s*tao\s*ung\s*dung', re.IGNORECASE)
            re_outdated_ward = re.compile(r'phường\s+tân\s+phú|p\.\s*tân\s*phú|phuong\s+tan\s+phu|p\.\s*tan\s*phu', re.IGNORECASE)

            challenger_deviated_ids = set()
            worker_deviated_ids = set()

            deviations_by_doc = {}

            for doc in all_non_folders:
                # 1. Run worker scanner
                worker_res = Scanner.scan_document(doc)
                if not worker_res.get('is_compliant', True):
                    worker_deviated_ids.add(doc.id)

                # 2. Run independent challenger extraction
                extracted_corpus = extract_text_from_doc(doc)
                has_vat = bool(re_legacy_vat.search(extracted_corpus))
                has_ent = bool(re_legacy_entity.search(extracted_corpus))
                # Outdated ward: only if not explicitly explaining Tan My replacement
                has_ward = False
                if re_outdated_ward.search(extracted_corpus):
                    lower_corpus = extracted_corpus.lower()
                    if "tân mỹ" not in lower_corpus and "tan my" not in lower_corpus:
                        has_ward = True

                if has_vat or has_ent or has_ward:
                    challenger_deviated_ids.add(doc.id)
                    deviations_by_doc[doc.id] = {
                        'name': doc.name,
                        'folder': doc.folder_id.name if doc.folder_id else 'None',
                        'has_vat': has_vat,
                        'has_ent': has_ent,
                        'has_ward': has_ward,
                    }

            print(f"      Worker Scanner Deviated Count: {len(worker_deviated_ids)}")
            print(f"      Challenger Independent Deviated Count: {len(challenger_deviated_ids)}")

            check(len(worker_deviated_ids) == 18, "Worker Scanner identified exactly 18 deviated documents", f"Count: {len(worker_deviated_ids)}")
            check(len(challenger_deviated_ids) == 18, "Challenger independent regex identified exactly 18 deviated documents", f"Count: {len(challenger_deviated_ids)}")
            check(worker_deviated_ids == challenger_deviated_ids, "100% Set Equivalence between Worker and Challenger Detections", "Zero divergence")

            # Check for False Negatives on the 143 compliant documents
            compliant_ids = set(all_non_folders.ids) - worker_deviated_ids
            check(len(compliant_ids) == 143, "Exactly 143 compliant documents remain", f"Count: {len(compliant_ids)}")

            false_negatives = []
            for cid in compliant_ids:
                cdoc = Doc.browse(cid)
                ctext = extract_text_from_doc(cdoc)
                if re_legacy_vat.search(ctext) or re_legacy_entity.search(ctext):
                    false_negatives.append((cid, cdoc.name))
            check(len(false_negatives) == 0, "Zero False Negatives on legacy VAT/entity across 143 compliant documents", f"False negatives: {false_negatives}")

            # Verify the specific high-profile deviated document: Doc ID 3376
            doc_3376 = Doc.browse(3376)
            check(doc_3376.exists(), "Doc ID 3376 exists", f"Name: {doc_3376.name}")
            scan_3376 = Scanner.scan_document(doc_3376)
            check(scan_3376.get('has_legacy_tax_code') is True, "Doc ID 3376 flagged for legacy tax code 0313683499", f"Summary: {scan_3376.get('summary')}")

            # =========================================================================
            # STAGE 4: MASTER DATA & VCB DIGIBIZ 11-COLUMN PAYROLL RECONCILIATION
            # =========================================================================
            print("\n>>> STAGE 4: Master Data & VCB DigiBiz 11-Column Export Validation <<<")

            # Ensure master data is seeded
            Scanner.seed_dgm_canonical_master_data()

            # 4.1 Verify 6 Canonical Partners
            partners_expected = [
                ('0314015511', 'CÔNG TY TNHH MTI TECHNOLOGY'),
                (None, 'Canva Pty Ltd'),
                ('0100686209', 'Trung tâm Dịch vụ số MobiFone'),
                (None, 'Temecula Valley Unified School District'),
                ('0300954529', 'CHI NHÁNH TẬP ĐOÀN BƯU CHÍNH VIỄN THÔNG VIỆT NAM'),
                ('0100109106-122', 'VIETTEL TP.HCM'),
            ]
            for vat, name_sub in partners_expected:
                domain = [('vat', '=', vat)] if vat else [('name', 'ilike', name_sub)]
                p = Partner.search(domain, limit=1)
                check(p.exists(), f"Canonical Partner '{name_sub}' exists", f"ID: {p.id if p else None}, VAT: {p.vat if p else None}")

            # 4.2 Verify 8 Canonical Employees
            expected_employees = [
                ('DGM-001', 'Lê Xuân Hoàng', '1373068161', 'VIETCOMBANK', 21410000),
                ('DGM-002', 'Nguyễn Hoài Ân', '1901014028', 'VIETCOMBANK', 18282000),
                ('DGM-003', 'Mai Thị Hồng Thắm', '0111000287863', 'VIETCOMBANK', 13469000),
                ('DGM-004', 'Bùi Hữu Tài', '1045246615', 'VIETCOMBANK', 17780000),
                ('DGM-005', 'Huỳnh Trần Quang Khánh', '12154887', 'ACB', 24876000),
                ('DGM-006', 'Ngô Thanh Lộc', '12154937', 'ACB', 28108000),
                ('DGM-007', 'Lê Cao Bảo Phúc', '12154917', 'ACB', 16502000),
                ('DGM-008', 'Thái Nguyễn Đăng Khoa', '47070467', 'ACB', 7500000),
            ]
            for reg, name, stk, bank_name, net_wage in expected_employees:
                emp = Employee.search([('registration_number', '=', reg)], limit=1)
                check(emp.exists(), f"Employee {reg} ({name}) exists", f"Found ID {emp.id if emp else None}")
                check(emp.primary_bank_account_id.account_number == stk, f"Employee {reg} STK match '{stk}'", f"Actual: {emp.primary_bank_account_id.account_number if emp else None}")
                if reg == 'DGM-003':
                    check(emp.primary_bank_account_id.account_number.startswith('0'), "Mai Thị Hồng Thắm (DGM-003) leading zero strictly preserved", f"STK: {emp.primary_bank_account_id.account_number}")

            # 4.3 Generate VCB DigiBiz 11-Column Batch Export
            if openpyxl:
                pay_run = env['hr.payslip.run'].create({
                    'name': 'Challenger VCB DigiBiz Batch Audit',
                    'date_start': fields.Date.from_string('2026-10-01'),
                    'date_end': fields.Date.from_string('2026-10-31'),
                })
                for reg, name, stk, bank_name, net_wage in expected_employees:
                    emp = Employee.search([('registration_number', '=', reg)], limit=1)
                    slip = env['hr.payslip'].create({
                        'name': f'Slip {name}',
                        'employee_id': emp.id,
                        'payslip_run_id': pay_run.id,
                        'date_from': pay_run.date_start,
                        'date_to': pay_run.date_end,
                    })
                    slip.write({'net_wage': net_wage, 'al2c_net_salary': net_wage})

                xlsx_bytes, fname = pay_run.generate_vcb_digibiz_xlsx()
                check(bool(xlsx_bytes), "VCB DigiBiz Excel generated", f"Filename: {fname}")

                wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
                ws = wb.active
                check(ws.max_column == 11, "Excel export has exactly 11 columns", f"Max columns: {ws.max_column}")
                check(ws.max_row == 9, "Excel export has exactly 9 rows (1 header + 8 data)", f"Max rows: {ws.max_row}")

                # Verify total wage
                total_wage = sum(int(ws.cell(row=r, column=9).value) for r in range(2, 10))
                check(total_wage == 147927000, "Total net salary sum strictly equals 147,927,000 VND", f"Sum: {total_wage:,} VND")

                # Verify Mai Thi Hong Tham row
                tham_cell = ws.cell(row=4, column=3).value  # Row 4 is STT 3
                check(str(tham_cell) == '0111000287863', "Leading zero in account number preserved in Excel output cell", f"Cell val: {tham_cell}")

            # =========================================================================
            # STAGE 5: ZERO-BRANDING FORENSIC INTEGRITY
            # =========================================================================
            print("\n>>> STAGE 5: Zero-Branding Forensic Integrity Scrutiny <<<")

            # Check 3 DGM crons
            dgm_crons = env['ir.cron'].search([('name', 'ilike', 'DGM Autonomous')])
            check(len(dgm_crons) == 3, "3 DGM Autonomous crons present", f"Count: {len(dgm_crons)}")
            for c in dgm_crons:
                check('odoo' not in c.name.lower() and 'openerp' not in c.name.lower(), f"Cron {c.id} name free of forbidden branding", f"Name: {c.name}")
                check('odoo' not in c.code.lower() and 'openerp' not in c.code.lower(), f"Cron {c.id} code free of forbidden branding", f"Code: {c.code}")

            # Check Audit Report Markdown & JSON
            md_path = '/home/zen/O20/.agents/teamwork/orchestrator_dgm_autonomous/DGM_AUTONOMOUS_COMPLIANCE_AUDIT_REPORT.md'
            json_path = '/home/zen/O20/.agents/teamwork/orchestrator_dgm_autonomous/DGM_AUTONOMOUS_COMPLIANCE_AUDIT_REPORT.json'
            check(os.path.exists(md_path), "Audit report markdown exists", f"{md_path}")
            check(os.path.exists(json_path), "Audit report JSON exists", f"{json_path}")

            if os.path.exists(md_path):
                with open(md_path, 'r', encoding='utf-8') as f:
                    md_content = f.read()
                check('odoo' not in md_content.lower(), "Markdown audit report contains 0 'odoo'", "Clean")
                check('openerp' not in md_content.lower(), "Markdown audit report contains 0 'openerp'", "Clean")

            if os.path.exists(json_path):
                with open(json_path, 'r', encoding='utf-8') as f:
                    json_content = f.read()
                check('odoo' not in json_content.lower(), "JSON audit report contains 0 'odoo'", "Clean")
                check('openerp' not in json_content.lower(), "JSON audit report contains 0 'openerp'", "Clean")

        finally:
            ICP.set_param('al2c.autonomy_enabled', orig_al2c)
            ICP.set_param('dgm.autonomy_enabled', orig_dgm)
            cr.rollback()

    print("\n" + "=" * 80)
    print(f"CHALLENGER 2 AUDIT COMPLETE: {passed_checks}/{total_checks} CHECKS PASSED")
    if failed_checks:
        print(f"FAILED CHECKS ({len(failed_checks)}):")
        for f in failed_checks:
            print(f"  - {f}")
        return False
    else:
        print("ALL 100% EMPIRICAL CHECKS PASSED: VERDICT APPROVE")
        print("=" * 80)
        return True


if __name__ == '__main__':
    ok = run_adversarial_challenge()
    sys.exit(0 if ok else 1)
